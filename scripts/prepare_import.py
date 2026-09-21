# Created by Harsh (@harsh-91) | Made in India | SPDX-License-Identifier: Apache-2.0
from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl


DATE_RE = re.compile(r"\d{2}/\d{2}/\d{4}")
PERIOD_RE = re.compile(r"Statement From\s*:\s*(\d{2}-\d{2}-\d{4})\s+to\s+(\d{2}-\d{2}-\d{4})")


def clean_decimal(value):
    if value is None or str(value).strip() == "":
        return None
    return Decimal(str(value).replace(",", "").strip())


def parse_labeled_value(value: str):
    if ":" not in value:
        return None
    key, raw = value.split(":", 1)
    key = re.sub(r"\s+", " ", key).strip().lower().replace(" ", "_")
    return key, raw.strip()


def read_statement(path: Path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    header_cells = []
    metadata = {}

    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=17, values_only=True), 1):
        for column_number, value in enumerate(row, 1):
            if value is None or str(value).strip() == "":
                continue
            text = str(value).strip()
            header_cells.append({"row": row_number, "column": column_number, "value": text})
            parsed = parse_labeled_value(text)
            if parsed:
                metadata[parsed[0]] = parsed[1]

    metadata["header_cells"] = header_cells
    statement_date = datetime.strptime(metadata["date_of_statement"], "%d-%m-%Y").date()
    period_match = PERIOD_RE.search(metadata["statement_from"] if "statement_from" in metadata else "")
    if period_match is None:
        period_text = next(cell["value"] for cell in header_cells if cell["value"].startswith("Statement From"))
        period_match = PERIOD_RE.search(period_text)
    if period_match is None:
        raise ValueError(f"Statement period was not found in {path.name}")

    period_start = datetime.strptime(period_match.group(1), "%d-%m-%Y").date()
    period_end = datetime.strptime(period_match.group(2), "%d-%m-%Y").date()
    account_number = metadata["account_number"]
    source_file = path.name.replace("-decrypted", "")
    metadata_json = json.dumps(metadata, ensure_ascii=False, separators=(",", ":"))
    records = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=19, values_only=True), 19):
        values = list(row[:6])
        raw_date = str(values[0] or "").strip()
        if not DATE_RE.fullmatch(raw_date):
            continue
        debit = clean_decimal(values[3])
        credit = clean_decimal(values[4])
        balance = clean_decimal(values[5])
        if (debit is None) == (credit is None):
            raise ValueError(f"Expected exactly one debit or credit in {path.name}, row {row_number}")
        if balance is None:
            raise ValueError(f"Balance is missing in {path.name}, row {row_number}")
        records.append(
            {
                "account_number": account_number,
                "statement_date": statement_date.isoformat(),
                "statement_period_start": period_start.isoformat(),
                "statement_period_end": period_end.isoformat(),
                "source_file": source_file,
                "source_row": row_number,
                "transaction_date": datetime.strptime(raw_date, "%d/%m/%Y").date().isoformat(),
                "details": str(values[1] or ""),
                "reference_number": str(values[2] or "") or None,
                "debit": debit,
                "credit": credit,
                "balance": balance,
                "statement_metadata": metadata_json,
            }
        )
    return records


def main():
    parser = argparse.ArgumentParser(description="Prepare SBI statement rows for PostgreSQL COPY.")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("files", nargs="+", type=Path)
    args = parser.parse_args()

    records = []
    summaries = []
    for path in args.files:
        file_records = read_statement(path)
        records.extend(file_records)
        summaries.append(
            {
                "file": path.name.replace("-decrypted", ""),
                "rows": len(file_records),
                "debits": sum(1 for row in file_records if row["debit"] is not None),
                "credits": sum(1 for row in file_records if row["credit"] is not None),
            }
        )

    fieldnames = list(records[0])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)

    print(json.dumps({"files": summaries, "total_rows": len(records), "output": str(args.output)}, indent=2))


if __name__ == "__main__":
    main()
