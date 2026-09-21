# Created by Harsh (@harsh-91) | Made in India | SPDX-License-Identifier: Apache-2.0
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

import msoffcrypto
import openpyxl


MONEY = Decimal("0.01")
MAX_WORKBOOK_FILES = 10_000
MAX_WORKBOOK_UNCOMPRESSED = 250 * 1024 * 1024
MAX_COMPRESSION_RATIO = 250


class StatementError(Exception):
    """A statement cannot be safely imported."""


class PasswordRequired(StatementError):
    """The workbook is encrypted and needs a password."""


@dataclass
class ParsedStatement:
    bank_key: str
    bank_name: str
    table_name: str
    account_number: str
    period_start: date
    period_end: date
    source_file: str
    source_sheet: str
    rows: list[dict[str, Any]]
    reconciliations: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def first_transaction_date(self) -> date:
        return min(row["transaction_date"] for row in self.rows)

    @property
    def last_transaction_date(self) -> date:
        return max(row["transaction_date"] for row in self.rows)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_text(value: Any) -> str:
    return re.sub(r"\s+", " ", clean_text(value)).strip()


def money(value: Any, *, allow_blank: bool = True) -> Decimal | None:
    raw = clean_text(value)
    if not raw:
        if allow_blank:
            return None
        raise StatementError("A required amount is blank")
    suffix = ""
    upper = raw.upper()
    if upper.endswith("CR") or upper.endswith("DR"):
        suffix = upper[-2:]
        raw = raw[:-2]
    raw = raw.replace(",", "").replace("₹", "").strip()
    try:
        result = Decimal(raw).quantize(MONEY)
    except InvalidOperation as error:
        raise StatementError(f"Invalid amount: {value!r}") from error
    return -result if suffix == "DR" else result


def date_value(value: Any, *formats: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = clean_text(value)
    for format_string in formats:
        try:
            return datetime.strptime(raw, format_string).date()
        except ValueError:
            continue
    raise StatementError(f"Invalid date: {value!r}")


def amount_token(value: Decimal | None) -> str:
    return "" if value is None else f"{value.quantize(MONEY):.2f}"


def fingerprint(bank_key: str, *parts: Any) -> str:
    values = [bank_key]
    for part in parts:
        if isinstance(part, Decimal):
            values.append(amount_token(part))
        elif isinstance(part, (date, datetime)):
            values.append(part.date().isoformat() if isinstance(part, datetime) else part.isoformat())
        else:
            values.append(normalized_text(part))
    return hashlib.sha256("|".join(values).encode("utf-8")).hexdigest()


def sbi_fingerprint(row: dict[str, Any]) -> str:
    return fingerprint(
        "sbi",
        row["account_number"], row["transaction_date"], row["details"],
        row.get("reference_number"), row.get("debit"), row.get("credit"), row["balance"],
    )


def icici_fingerprint(row: dict[str, Any]) -> str:
    return fingerprint(
        "icici",
        row["account_number"], row["value_date"], row["transaction_date"],
        row.get("cheque_number"), row["transaction_remarks"],
        row["withdrawal_amount"], row["deposit_amount"], row["balance"],
    )


def indusind_fingerprint(row: dict[str, Any]) -> str:
    return fingerprint(
        "indusind",
        row["account_number"], row["transaction_date"], row["transaction_type"],
        row["description"], row.get("debit"), row.get("credit"), row["balance"],
    )


def _load_workbook(data: bytes, password: str | None):
    if data.startswith(b"PK\x03\x04"):
        plain = data
    elif data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        decrypted = BytesIO()
        try:
            office = msoffcrypto.OfficeFile(BytesIO(data))
            if not password:
                raise PasswordRequired("This statement is password protected")
            office.load_key(password=password)
            office.decrypt(decrypted)
            plain = decrypted.getvalue()
        except PasswordRequired:
            raise
        except Exception as error:
            raise StatementError("Could not decrypt the statement. Check the password.") from error
    else:
        raise StatementError("Unsupported file format. Upload an Excel .xlsx or bank-exported .xls file.")
    _validate_ooxml_archive(plain)
    try:
        return openpyxl.load_workbook(BytesIO(plain), read_only=True, data_only=True)
    except Exception as error:
        raise StatementError("The workbook could not be read as an Excel statement.") from error


def _validate_ooxml_archive(data: bytes) -> None:
    """Reject malformed and suspiciously expanded OOXML archives before parsing."""
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_WORKBOOK_FILES:
                raise StatementError("The workbook contains too many internal files")
            total = sum(item.file_size for item in members)
            if total > MAX_WORKBOOK_UNCOMPRESSED:
                raise StatementError("The workbook expands beyond the 250 MB safety limit")
            for item in members:
                if item.file_size and item.compress_size == 0:
                    raise StatementError("The workbook contains an invalid compressed entry")
                if item.compress_size and item.file_size / item.compress_size > MAX_COMPRESSION_RATIO:
                    raise StatementError("The workbook contains a suspicious compression ratio")
    except StatementError:
        raise
    except (zipfile.BadZipFile, OSError) as error:
        raise StatementError("The workbook is not a valid Excel OOXML archive") from error


def _assert_period(rows: list[dict[str, Any]], period_start: date, period_end: date) -> None:
    if period_start > period_end:
        raise StatementError("Statement period start is after its end")
    outside = [row["source_row"] for row in rows if not period_start <= row["transaction_date"] <= period_end]
    if outside:
        raise StatementError(f"Transaction date falls outside the statement period at source row {outside[0]}")


def _assert_nonnegative_amounts(rows: list[dict[str, Any]], *keys: str) -> None:
    for row in rows:
        for key in keys:
            value = row.get(key)
            if value is not None and value < 0:
                raise StatementError(f"Negative {key.replace('_', ' ')} at source row {row['source_row']}")


def _cell_matrix(sheet, max_rows: int = 30, max_cols: int = 12) -> list[list[str]]:
    return [
        [normalized_text(value) for value in row[:max_cols]]
        for row in sheet.iter_rows(min_row=1, max_row=min(max_rows, sheet.max_row), values_only=True)
    ]


def detect_bank(workbook) -> str:
    sheet = workbook.active
    matrix = _cell_matrix(sheet)
    flattened = "\n".join("|".join(row) for row in matrix).lower()
    if "ref no/cheque no" in flattened and "state bank of india" in flattened:
        return "sbi"
    if "withdrawal amount(inr)" in flattened and "deposit amount(inr)" in flattened:
        return "icici"
    if "account information" in flattened and "transaction list" in flattened:
        for row in matrix:
            if row and row[0].lower().startswith("ifsc code") and len(row) > 1 and row[1].upper().startswith("INDB"):
                return "indusind"
    raise StatementError("Bank structure was not recognized. Supported banks: SBI, ICICI Bank, and IndusInd Bank.")


def _header_cells(sheet, last_row: int) -> list[dict[str, Any]]:
    result = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=last_row, values_only=True), 1):
        for column_number, value in enumerate(row, 1):
            if normalized_text(value):
                result.append({"row": row_number, "column": column_number, "value": clean_text(value)})
    return result


def _parse_labeled_headers(sheet, last_row: int) -> dict[str, str]:
    result: dict[str, str] = {}
    for cell in _header_cells(sheet, last_row):
        value = cell["value"]
        if ":" not in value:
            continue
        key, raw = value.split(":", 1)
        normalized_key = re.sub(r"\W+", "_", key).strip("_").lower()
        result[normalized_key] = raw.strip()
    return result


def _assert_balance_chain(rows: list[dict[str, Any]], debit_key: str, credit_key: str, *, reverse: bool = False) -> None:
    ordered = list(reversed(rows)) if reverse else rows
    mismatches = []
    for previous, current in zip(ordered, ordered[1:]):
        debit = current.get(debit_key) or Decimal(0)
        credit = current.get(credit_key) or Decimal(0)
        expected = (previous["balance"] + credit - debit).quantize(MONEY)
        if expected != current["balance"].quantize(MONEY):
            mismatches.append((current["source_row"], expected, current["balance"]))
    if mismatches:
        first = mismatches[0]
        raise StatementError(
            f"Running balance failed at source row {first[0]}: expected {first[1]}, found {first[2]} "
            f"({len(mismatches)} mismatch{'es' if len(mismatches) != 1 else ''})."
        )


def _reconcile_icici_balances(rows: list[dict[str, Any]]) -> list[str]:
    """Reconcile ICICI balances, including a bank-export reversal display anomaly.

    Some ICICI exports show an adjacent equal credit/debit reversal pair without
    applying either amount to the displayed intermediate balances. We accept
    that pair only when both remarks share the same long bank reference and the
    next transaction's displayed balance anchors the net-zero result.
    """
    logical_balance = rows[0]["balance"]
    reversal_pairs = 0
    index = 1
    while index < len(rows):
        current = rows[index]
        expected = (
            logical_balance
            + current["deposit_amount"]
            - current["withdrawal_amount"]
        ).quantize(MONEY)
        if expected == current["balance"].quantize(MONEY):
            logical_balance = current["balance"]
            index += 1
            continue

        if index + 2 < len(rows):
            reversal = rows[index + 1]
            anchor = rows[index + 2]
            current_amount = max(current["deposit_amount"], current["withdrawal_amount"])
            reversal_amount = max(reversal["deposit_amount"], reversal["withdrawal_amount"])
            opposite_sides = (
                current["deposit_amount"] == reversal["withdrawal_amount"]
                and current["withdrawal_amount"] == reversal["deposit_amount"]
                and current_amount > 0
            )
            current_references = set(re.findall(r"\d{10,}", current["transaction_remarks"]))
            reversal_references = set(re.findall(r"\d{10,}", reversal["transaction_remarks"]))
            shared_reference = bool(current_references & reversal_references)
            anchored_balance = (
                logical_balance
                + current["deposit_amount"] - current["withdrawal_amount"]
                + reversal["deposit_amount"] - reversal["withdrawal_amount"]
                + anchor["deposit_amount"] - anchor["withdrawal_amount"]
            ).quantize(MONEY)
            if opposite_sides and current_amount == reversal_amount and shared_reference and anchored_balance == anchor["balance"]:
                logical_balance = (logical_balance + current["deposit_amount"] - current["withdrawal_amount"]
                                   + reversal["deposit_amount"] - reversal["withdrawal_amount"]).quantize(MONEY)
                reversal_pairs += 1
                index += 2
                continue

        raise StatementError(
            f"Running balance failed at source row {current['source_row']}: "
            f"expected {expected}, found {current['balance']}."
        )
    messages = [f"Running balances reconcile across {len(rows)} transactions"]
    if reversal_pairs:
        messages.append(
            f"Verified {reversal_pairs} adjacent equal credit/debit reversal pair using its shared bank reference and next balance"
        )
    return messages


def parse_sbi(sheet, source_file: str) -> ParsedStatement:
    headers = _parse_labeled_headers(sheet, 17)
    full_headers = _header_cells(sheet, 17)
    account_number = headers.get("account_number")
    statement_date_raw = headers.get("date_of_statement")
    period_text = next((cell["value"] for cell in full_headers if cell["value"].startswith("Statement From")), "")
    period_match = re.search(r"(\d{2}-\d{2}-\d{4})\s+to\s+(\d{2}-\d{2}-\d{4})", period_text)
    if not account_number or not statement_date_raw or not period_match:
        raise StatementError("SBI statement metadata is incomplete")
    statement_date = date_value(statement_date_raw, "%d-%m-%Y")
    period_start = date_value(period_match.group(1), "%d-%m-%Y")
    period_end = date_value(period_match.group(2), "%d-%m-%Y")
    metadata = {"bank": "State Bank of India", "fields": headers, "header_cells": full_headers}
    metadata_json = json.dumps(metadata, ensure_ascii=False, separators=(",", ":"))
    rows: list[dict[str, Any]] = []
    summary_header_row = None
    for source_row, row in enumerate(sheet.iter_rows(min_row=19, values_only=True), 19):
        values = list(row[:6])
        if clean_text(values[0]).startswith("Brought Forward"):
            summary_header_row = source_row
            continue
        raw_date = clean_text(values[0])
        if not re.fullmatch(r"\d{2}/\d{2}/\d{4}", raw_date):
            continue
        parsed = {
            "account_number": account_number,
            "statement_date": statement_date,
            "statement_period_start": period_start,
            "statement_period_end": period_end,
            "source_file": source_file,
            "source_row": source_row,
            "transaction_date": date_value(raw_date, "%d/%m/%Y"),
            "details": clean_text(values[1]),
            "reference_number": clean_text(values[2]) or None,
            "debit": money(values[3]),
            "credit": money(values[4]),
            "balance": money(values[5], allow_blank=False),
            "statement_metadata": metadata_json,
        }
        if (parsed["debit"] is None) == (parsed["credit"] is None):
            raise StatementError(f"SBI source row {source_row} must contain exactly one debit or credit")
        parsed["transaction_fingerprint"] = sbi_fingerprint(parsed)
        rows.append(parsed)
    if not rows:
        raise StatementError("No SBI transactions were found")
    _assert_period(rows, period_start, period_end)
    _assert_nonnegative_amounts(rows, "debit", "credit")
    _assert_balance_chain(rows, "debit", "credit")
    reconciliations = [f"Running balances reconcile across {len(rows)} transactions"]
    if summary_header_row:
        summary = [sheet.cell(summary_header_row + 1, column).value for column in range(1, 7)]
        brought_forward = money(summary[0], allow_blank=False)
        debit_count = int(clean_text(summary[1]).replace(",", ""))
        credit_count = int(clean_text(summary[2]).replace(",", ""))
        debit_total = money(summary[3], allow_blank=False)
        credit_total = money(summary[4], allow_blank=False)
        closing_balance = money(summary[5], allow_blank=False)
        checks = {
            "debit count": (sum(row["debit"] is not None for row in rows), debit_count),
            "credit count": (sum(row["credit"] is not None for row in rows), credit_count),
            "debit total": (sum((row["debit"] or Decimal(0)) for row in rows), debit_total),
            "credit total": (sum((row["credit"] or Decimal(0)) for row in rows), credit_total),
            "closing balance": (rows[-1]["balance"], closing_balance),
            "opening balance": (
                (rows[0]["balance"] - (rows[0]["credit"] or Decimal(0)) + (rows[0]["debit"] or Decimal(0))).quantize(MONEY),
                brought_forward,
            ),
        }
        failed = [name for name, (actual, expected) in checks.items() if actual != expected]
        if failed:
            raise StatementError(f"SBI statement controls failed: {', '.join(failed)}")
        reconciliations.append("SBI counts, totals, opening balance, and closing balance match the statement controls")
    else:
        raise StatementError("SBI statement control totals were not found")
    return ParsedStatement("sbi", "State Bank of India", "bank_sbi_transactions", account_number, period_start, period_end, source_file, sheet.title, rows, reconciliations)


def parse_icici(sheet, source_file: str) -> ParsedStatement:
    account_display = clean_text(sheet.cell(4, 4).value)
    account_number = account_display.split()[0]
    period_start = date_value(sheet.cell(5, 4).value, "%d,%m,%Y", "%d/%m/%Y")
    period_end = date_value(sheet.cell(5, 6).value, "%d,%m,%Y", "%d/%m/%Y")
    metadata = {
        "bank": "ICICI Bank", "account_display": account_display,
        "statement_period_start": period_start.isoformat(), "statement_period_end": period_end.isoformat(),
        "header_cells": _header_cells(sheet, 13),
    }
    metadata_json = json.dumps(metadata, ensure_ascii=False, separators=(",", ":"))
    rows: list[dict[str, Any]] = []
    for source_row, row in enumerate(sheet.iter_rows(min_row=14, values_only=True), 14):
        values = list(row)
        serial = clean_text(values[1] if len(values) > 1 else None)
        if not serial.isdigit():
            continue
        parsed = {
            "account_number": account_number,
            "statement_period_start": period_start,
            "statement_period_end": period_end,
            "source_file": source_file,
            "source_sheet": sheet.title,
            "source_row": source_row,
            "serial_number": int(serial),
            "value_date": date_value(values[2], "%d,%m,%Y", "%d/%m/%Y"),
            "transaction_date": date_value(values[3], "%d,%m,%Y", "%d/%m/%Y"),
            "cheque_number": clean_text(values[4]) or None,
            "transaction_remarks": clean_text(values[5]),
            "withdrawal_amount": money(values[6], allow_blank=False),
            "deposit_amount": money(values[7], allow_blank=False),
            "balance": money(values[8], allow_blank=False),
            "statement_metadata": metadata_json,
        }
        parsed["transaction_fingerprint"] = icici_fingerprint(parsed)
        rows.append(parsed)
    if not rows:
        raise StatementError("No ICICI Bank transactions were found")
    _assert_period(rows, period_start, period_end)
    _assert_nonnegative_amounts(rows, "withdrawal_amount", "deposit_amount")
    reconciliations = _reconcile_icici_balances(rows)
    warnings = []
    if len(reconciliations) > 1:
        warnings.append("ICICI's displayed balances skip the intermediate effect of a matched reversal pair; the next balance reconciles its net-zero effect.")
    return ParsedStatement("icici", "ICICI Bank", "bank_icici_transactions", account_number, period_start, period_end, source_file, sheet.title, rows, reconciliations, warnings)


def parse_indusind(sheet, source_file: str) -> ParsedStatement:
    fields: dict[str, str] = {}
    for row_number in range(5, 14):
        key = normalized_text(sheet.cell(row_number, 1).value)
        value = clean_text(sheet.cell(row_number, 2).value)
        if key:
            fields[re.sub(r"\W+", "_", key).strip("_").lower()] = value
    account_number = fields.get("account_number", "")
    period_match = re.fullmatch(r"(\d{2} [A-Za-z]+ \d{4}) TO (\d{2} [A-Za-z]+ \d{4})", fields.get("transaction_from", ""))
    if not account_number or period_match is None or not fields.get("ifsc_code", "").upper().startswith("INDB"):
        raise StatementError("IndusInd statement metadata is incomplete")
    period_start = date_value(period_match.group(1), "%d %B %Y")
    period_end = date_value(period_match.group(2), "%d %B %Y")
    metadata = {"bank": "IndusInd Bank", "fields": fields, "header_cells": _header_cells(sheet, 22)}
    metadata_json = json.dumps(metadata, ensure_ascii=False, separators=(",", ":"))
    rows: list[dict[str, Any]] = []
    for source_row, row in enumerate(sheet.iter_rows(min_row=23, values_only=True), 23):
        values = list(row)
        serial = clean_text(values[0] if values else None)
        if not serial.isdigit():
            continue
        parsed = {
            "account_number": account_number,
            "statement_period_start": period_start,
            "statement_period_end": period_end,
            "source_file": source_file,
            "source_sheet": sheet.title,
            "source_row": source_row,
            "serial_number": int(serial),
            "transaction_date": date_value(values[1], "%Y-%m-%d", "%d/%m/%Y"),
            "transaction_type": clean_text(values[2]),
            "description": clean_text(values[3]),
            "debit": money(values[4]),
            "credit": money(values[5]),
            "balance": money(values[6], allow_blank=False),
            "statement_metadata": metadata_json,
        }
        if (parsed["debit"] is None) == (parsed["credit"] is None):
            raise StatementError(f"IndusInd source row {source_row} must contain exactly one debit or credit")
        parsed["transaction_fingerprint"] = indusind_fingerprint(parsed)
        rows.append(parsed)
    if not rows:
        raise StatementError("No IndusInd Bank transactions were found")
    _assert_period(rows, period_start, period_end)
    _assert_nonnegative_amounts(rows, "debit", "credit")
    _assert_balance_chain(rows, "debit", "credit", reverse=True)
    reconciliations = [f"Running balances reconcile across {len(rows)} transactions"]
    return ParsedStatement("indusind", "IndusInd Bank", "bank_indusind_transactions", account_number, period_start, period_end, source_file, sheet.title, rows, reconciliations)


def parse_statement(data: bytes, source_file: str, password: str | None = None) -> ParsedStatement:
    source_file = source_file.replace("\\", "/").rsplit("/", 1)[-1]
    workbook = _load_workbook(data, password)
    try:
        bank_key = detect_bank(workbook)
        sheet = workbook.active
        if bank_key == "sbi":
            return parse_sbi(sheet, source_file)
        if bank_key == "icici":
            return parse_icici(sheet, source_file)
        if bank_key == "indusind":
            return parse_indusind(sheet, source_file)
        raise StatementError("No parser is available for this bank")
    finally:
        workbook.close()
