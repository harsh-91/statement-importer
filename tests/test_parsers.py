# Created by Harsh (@harsh-91) | Made in India | SPDX-License-Identifier: Apache-2.0
import unittest
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
from msoffcrypto.format.ooxml import OOXMLFile

from statement_importer.parsers import PasswordRequired, StatementError, _load_workbook, date_value, fingerprint, money, parse_statement


class ParserPrimitiveTests(unittest.TestCase):
    def test_money_is_exact_and_locale_tolerant(self):
        self.assertEqual(money("₹1,234.50"), Decimal("1234.50"))
        self.assertEqual(money("100.00 DR"), Decimal("-100.00"))
        self.assertIsNone(money(""))

    def test_supported_dates_are_unambiguous(self):
        self.assertEqual(date_value("21/09/2026", "%d/%m/%Y"), date(2026, 9, 21))
        self.assertEqual(date_value("2026-09-21", "%Y-%m-%d"), date(2026, 9, 21))

    def test_fingerprint_is_stable(self):
        parts = ("ACCOUNT", date(2026, 9, 21), "Coffee", Decimal("125.00"), Decimal("875.00"))
        self.assertEqual(fingerprint("sample_bank", *parts), fingerprint("sample_bank", *parts))
        self.assertNotEqual(fingerprint("sample_bank", *parts), fingerprint("another_bank", *parts))

    def test_plain_icici_xls_imports_without_password(self):
        path = Path(__file__).parent / "fixtures" / "icici-plain.xls"
        for password in (None, "unneeded"):
            with self.subTest(password=password):
                statement = parse_statement(path.read_bytes(), path.name, password)
                self.assertEqual(statement.bank_name, "ICICI Bank")
                self.assertEqual(len(statement.rows), 1)
                self.assertEqual(statement.rows[0]["transaction_date"], date(2026, 9, 2))
                self.assertEqual(statement.rows[0]["balance"], Decimal("100.00"))

    def test_encrypted_office_file_still_requests_password(self):
        plain = BytesIO()
        workbook = openpyxl.Workbook()
        workbook.save(plain)
        encrypted = BytesIO()
        OOXMLFile(BytesIO(plain.getvalue())).encrypt("test-password", encrypted)
        with self.assertRaises(PasswordRequired):
            _load_workbook(encrypted.getvalue(), None)
        opened = _load_workbook(encrypted.getvalue(), "test-password")
        opened.close()

    def test_corrupt_ole_file_is_not_reported_as_password_protected(self):
        with self.assertRaises(StatementError) as caught:
            _load_workbook(bytes.fromhex("D0CF11E0A1B11AE1") + b"broken", None)
        self.assertNotIsInstance(caught.exception, PasswordRequired)


if __name__ == "__main__":
    unittest.main()
